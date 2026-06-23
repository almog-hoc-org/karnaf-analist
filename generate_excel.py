#!/usr/bin/env python3
"""Generate comprehensive Excel report from Israeli real estate database."""

import sqlite3
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

# Paths
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "realestate.db")
OUTPUT_PATH = os.path.join(BASE_DIR, "data", "realestate_data.xlsx")

# Style constants
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Arial", size=10)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
RTL_ALIGN = Alignment(horizontal="right", readingOrder=2)
RTL_CENTER = Alignment(horizontal="center", readingOrder=2)
NUM_ALIGN = Alignment(horizontal="center", readingOrder=2)
THIN_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
ALT_FILL = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
SUM_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
SUM_FONT = Font(name="Arial", bold=True, size=10)

THOUSANDS_FMT = '#,##0'
DECIMAL_FMT = '#,##0.00'
PCT_FMT = '0.00%'
PCT_PLAIN_FMT = '#,##0.00"%"'


def apply_header_style(ws, row=1, max_col=None):
    """Apply header styling to first row."""
    if max_col is None:
        max_col = ws.max_column
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def apply_data_style(ws, start_row=2, max_row=None, max_col=None, num_cols=None, pct_cols=None, decimal_cols=None, text_cols=None):
    """Apply formatting to data cells."""
    if max_row is None:
        max_row = ws.max_row
    if max_col is None:
        max_col = ws.max_column
    if num_cols is None:
        num_cols = set()
    if pct_cols is None:
        pct_cols = set()
    if decimal_cols is None:
        decimal_cols = set()
    if text_cols is None:
        text_cols = set()

    for row in range(start_row, max_row + 1):
        is_alt = (row - start_row) % 2 == 1
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if is_alt:
                cell.fill = ALT_FILL
            if col in text_cols:
                cell.alignment = RTL_ALIGN
            elif col in pct_cols:
                cell.alignment = NUM_ALIGN
                cell.number_format = PCT_PLAIN_FMT
            elif col in decimal_cols:
                cell.alignment = NUM_ALIGN
                cell.number_format = DECIMAL_FMT
            elif col in num_cols:
                cell.alignment = NUM_ALIGN
                cell.number_format = THOUSANDS_FMT
            else:
                cell.alignment = RTL_CENTER


def set_column_widths(ws, widths):
    """Set column widths from dict {col_num: width}."""
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def setup_sheet_rtl(ws):
    """Set sheet to RTL and freeze header."""
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"


def main():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    wb = Workbook()

    # ===== SHEET 1: Cities Base Data =====
    ws1 = wb.active
    ws1.title = "ערים - נתוני יסוד"
    setup_sheet_rtl(ws1)

    headers1 = [
        "עיר", "אוכלוסייה 2021", "אוכלוסייה 2022", "אוכלוסייה 2024",
        "אוכלוסייה 2026", "משקי בית 2022", "גודל משק בית ממוצע",
        "נפשות לדירה", "בנייה גולמית 4 שנים", "מקדם נטו",
        "מחיר למ\"ר 2023", "מחיר למ\"ר 2026", "שינוי מחיר %",
        "סה\"כ דירות", "בנייה נטו", "גידול אוכלוסייה מוחלט",
        "גידול אוכלוסייה %", "דירות נדרשות", "גידול דירות",
        "מכפיל מוזהב", "אחוז מוזהב %", "יחידות חידוש עירוני קיימות",
        "יחידות חידוש עירוני מוצעות", "סטטוס חידוש עירוני"
    ]
    ws1.append(headers1)
    apply_header_style(ws1, max_col=len(headers1))

    cur.execute("""
        SELECT city_name, population_2021, population_2022, population_2024,
               population_2026, households_2022, avg_household_size_2022,
               people_per_apartment, construction_4y_gross, net_coefficient,
               price_per_sqm_2023, price_per_sqm_2026, price_change_pct,
               total_apartments, construction_net, population_growth_abs,
               population_growth_pct, apartments_required, apartment_growth,
               golden_multiplier, golden_pct, urban_renewal_existing_units,
               urban_renewal_proposed_units, urban_renewal_status
        FROM cities
        ORDER BY COALESCE(population_2026, population_2024, population_2022, population_2021, 0) DESC
    """)
    for row in cur.fetchall():
        ws1.append(list(row))

    num_cols_1 = {2,3,4,5,6,9,11,12,14,15,16,22,23}
    decimal_cols_1 = {7,8,10,18,19,20}
    pct_cols_1 = {13,17,21}
    text_cols_1 = {1,24}
    apply_data_style(ws1, num_cols=num_cols_1, pct_cols=pct_cols_1, decimal_cols=decimal_cols_1, text_cols=text_cols_1)

    widths1 = {1:18, 2:14, 3:14, 4:14, 5:14, 6:14, 7:14, 8:12, 9:16, 10:12,
               11:14, 12:14, 13:12, 14:12, 15:12, 16:16, 17:14, 18:12, 19:12,
               20:12, 21:12, 22:18, 23:18, 24:18}
    set_column_widths(ws1, widths1)
    ws1.auto_filter.ref = ws1.dimensions

    # ===== SHEET 2: Building Permits Pivot =====
    ws2 = wb.create_sheet("היתרי בנייה לפי שנה")
    setup_sheet_rtl(ws2)

    years = list(range(2016, 2025))
    headers2 = ["עיר"] + [str(y) for y in years] + ["סה\"כ", "ממוצע"]
    ws2.append(headers2)
    apply_header_style(ws2, max_col=len(headers2))

    # Get all permits pivoted
    cur.execute("""
        SELECT c.city_name,
            MAX(CASE WHEN bp.year=2016 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2017 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2018 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2019 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2020 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2021 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2022 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2023 THEN bp.permits END),
            MAX(CASE WHEN bp.year=2024 THEN bp.permits END)
        FROM cities c
        LEFT JOIN building_permits bp ON bp.city_name = c.city_name
        GROUP BY c.city_name
        ORDER BY c.city_name
    """)
    rows2 = cur.fetchall()
    for i, row in enumerate(rows2):
        data_row = i + 2  # 1-based, row 1 is header
        ws2.append(list(row))
        # Add SUM formula (columns B=2 through J=10)
        sum_col = len(years) + 2  # column 11
        avg_col = sum_col + 1      # column 12
        b_letter = get_column_letter(2)
        j_letter = get_column_letter(len(years) + 1)
        ws2.cell(row=data_row, column=sum_col, value=f"=SUM({b_letter}{data_row}:{j_letter}{data_row})")
        ws2.cell(row=data_row, column=avg_col, value=f"=AVERAGE({b_letter}{data_row}:{j_letter}{data_row})")

    num_cols_2 = set(range(2, len(headers2) + 1))
    text_cols_2 = {1}
    apply_data_style(ws2, num_cols=num_cols_2, text_cols=text_cols_2)

    # Style SUM/AVG columns
    for row in range(2, ws2.max_row + 1):
        for col in [len(years) + 2, len(years) + 3]:
            cell = ws2.cell(row=row, column=col)
            cell.font = SUM_FONT
            cell.number_format = THOUSANDS_FMT

    widths2 = {1: 18}
    for c in range(2, len(headers2) + 1):
        widths2[c] = 12
    set_column_widths(ws2, widths2)
    ws2.auto_filter.ref = ws2.dimensions

    # ===== SHEET 3: Sales Data =====
    ws3 = wb.create_sheet("מכירות חדשות")
    setup_sheet_rtl(ws3)

    headers3 = [
        "עיר", "מכירות חדשות 2023", "מכירות חדשות 2024", "מכירות חדשות 2025",
        "מלאי לא מכור 2025", "ממוצע מכירות 3 שנים", "שנים לפינוי 2025",
        "שנים לפינוי ממוצע"
    ]
    ws3.append(headers3)
    apply_header_style(ws3, max_col=len(headers3))

    cur.execute("""
        SELECT city_name, new_sales_2023, new_sales_2024, new_sales_2025,
               unsold_inventory_2025, avg_sales_3y, years_to_clear_2025, years_to_clear_avg
        FROM city_sales
        ORDER BY city_name
    """)
    for row in cur.fetchall():
        ws3.append(list(row))

    num_cols_3 = {2,3,4,5}
    decimal_cols_3 = {6,7,8}
    text_cols_3 = {1}
    apply_data_style(ws3, num_cols=num_cols_3, decimal_cols=decimal_cols_3, text_cols=text_cols_3)

    widths3 = {1:18, 2:16, 3:16, 4:16, 5:16, 6:18, 7:16, 8:16}
    set_column_widths(ws3, widths3)
    ws3.auto_filter.ref = ws3.dimensions

    # ===== SHEET 4: Population Trends =====
    ws4 = wb.create_sheet("אוכלוסייה - מגמות")
    setup_sheet_rtl(ws4)

    headers4 = [
        "עיר", "אוכלוסייה 2021", "אוכלוסייה 2022", "אוכלוסייה 2024",
        "אוכלוסייה 2026", "גידול 2022-2024", "שיעור גידול שנתי %"
    ]
    ws4.append(headers4)
    apply_header_style(ws4, max_col=len(headers4))

    # Only cities with at least 2 population data points
    cur.execute("""
        SELECT city_name, population_2021, population_2022, population_2024, population_2026
        FROM cities
        WHERE (CASE WHEN population_2021 IS NOT NULL THEN 1 ELSE 0 END +
               CASE WHEN population_2022 IS NOT NULL THEN 1 ELSE 0 END +
               CASE WHEN population_2024 IS NOT NULL THEN 1 ELSE 0 END +
               CASE WHEN population_2026 IS NOT NULL THEN 1 ELSE 0 END) >= 2
        ORDER BY COALESCE(population_2026, population_2024, population_2022, population_2021, 0) DESC
    """)
    rows4 = cur.fetchall()
    for i, row in enumerate(rows4):
        data_row = i + 2
        ws4.append(list(row))
        # Growth 2022 vs 2024: =D{r}-C{r}
        ws4.cell(row=data_row, column=6,
                 value=f'=IF(AND(D{data_row}<>"",C{data_row}<>""),D{data_row}-C{data_row},"")')
        # Annual growth rate: if we have 2022 and 2024, = (D/C)^(1/2)-1
        # If we have 2021 and 2026: = (E/B)^(1/5)-1
        # Use best available span
        ws4.cell(row=data_row, column=7,
                 value=f'=IF(AND(D{data_row}<>"",C{data_row}<>"",C{data_row}<>0),((D{data_row}/C{data_row})^(1/2)-1)*100,IF(AND(E{data_row}<>"",B{data_row}<>"",B{data_row}<>0),((E{data_row}/B{data_row})^(1/5)-1)*100,""))')

    num_cols_4 = {2,3,4,5,6}
    decimal_cols_4 = {7}
    text_cols_4 = {1}
    apply_data_style(ws4, num_cols=num_cols_4, decimal_cols=decimal_cols_4, text_cols=text_cols_4)

    widths4 = {1:18, 2:14, 3:14, 4:14, 5:14, 6:16, 7:18}
    set_column_widths(ws4, widths4)
    ws4.auto_filter.ref = ws4.dimensions

    # Save
    wb.save(OUTPUT_PATH)
    conn.close()
    print(f"Excel file created successfully: {OUTPUT_PATH}")
    print(f"  Sheet 1 (ערים - נתוני יסוד): {ws1.max_row - 1} cities")
    print(f"  Sheet 2 (היתרי בנייה לפי שנה): {ws2.max_row - 1} cities x {len(years)} years")
    print(f"  Sheet 3 (מכירות חדשות): {ws3.max_row - 1} rows")
    print(f"  Sheet 4 (אוכלוסייה - מגמות): {ws4.max_row - 1} cities")


if __name__ == "__main__":
    main()
